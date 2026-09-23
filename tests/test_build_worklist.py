"""The opt-out worklist. It IS the tracker, so rebuilding must never lose work."""
from __future__ import annotations

import pytest
import yaml
from openpyxl import Workbook, load_workbook

import build_worklist as BW
from client_context import ClientDataError, output_dir

SITES = [
    {"key": "whitepages", "name": "Whitepages", "category": "data_broker",
     "opt_out_url": "https://example.invalid/wp", "needs": ["full_name", "home_city"],
     "supports_agent": False, "notes": "Locate the profile URL first."},
    {"key": "spokeo", "name": "Spokeo", "category": "data_broker",
     "opt_out_url": "https://example.invalid/sp", "needs": ["full_name", "email"],
     "supports_agent": False, "notes": "Needs the exact profile URL."},
]


@pytest.fixture
def sites_file(tmp_path):
    path = tmp_path / "sites.yaml"
    path.write_text(yaml.safe_dump(SITES), encoding="utf-8")
    return path


def read_rows(path):
    """{site name: {column: value}} from a built worklist."""
    ws = load_workbook(path).active
    headers = [c.value for c in ws[BW.HEADER_ROW]]
    rows = {}
    for row in ws.iter_rows(min_row=BW.HEADER_ROW + 1, values_only=True):
        if row and row[0]:
            rows[row[0]] = dict(zip(headers, row))
    return rows


def set_cell(path, site, column, value):
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[BW.HEADER_ROW]]
    col = headers.index(column) + 1
    for r in range(BW.HEADER_ROW + 1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == site:
            ws.cell(row=r, column=col, value=value)
    wb.save(path)


class TestBuild:
    def test_writes_one_row_per_site(self, saved_client, sites_file):
        path = BW.build_worklist(saved_client, sites_path=sites_file)
        assert set(read_rows(path)) == {"Whitepages", "Spokeo"}

    def test_prefills_the_values_each_site_asks_for(self, saved_client, sites_file):
        path = BW.build_worklist(saved_client, sites_path=sites_file)
        prefilled = read_rows(path)["Whitepages"]["Prefilled values (copy/paste)"]
        assert "Alexandra M Example" in prefilled and "Brandon" in prefilled

    def test_starts_every_row_as_not_started(self, saved_client, sites_file):
        path = BW.build_worklist(saved_client, sites_path=sites_file)
        assert read_rows(path)["Spokeo"]["Status"] == "not_started"

    def test_lands_in_the_client_s_own_folder(self, saved_client, sites_file):
        path = BW.build_worklist(saved_client, sites_path=sites_file)
        assert path.parent == output_dir(saved_client)

    def test_catalog_guidance_is_its_own_column(self, saved_client, sites_file):
        path = BW.build_worklist(saved_client, sites_path=sites_file)
        row = read_rows(path)["Whitepages"]
        assert row["Site guidance"] == "Locate the profile URL first."
        assert not row["Notes"]              # the worker's column starts empty


class TestRebuildPreservesWork:
    @pytest.fixture
    def worked(self, saved_client, sites_file):
        path = BW.build_worklist(saved_client, sites_path=sites_file)
        set_cell(path, "Whitepages", "Status", "removed")
        set_cell(path, "Whitepages", "Submitted", "2026-09-01")
        set_cell(path, "Whitepages", "Confirmation / ref", "WP-88213")
        set_cell(path, "Whitepages", "Notes", "phone verification, two tries")
        set_cell(path, "Spokeo", "Status", "verification_sent")
        return path

    def test_keeps_the_status(self, saved_client, sites_file, worked):
        BW.build_worklist(saved_client, sites_path=sites_file)
        assert read_rows(worked)["Whitepages"]["Status"] == "removed"

    def test_keeps_the_confirmation_and_the_worker_s_notes(self, saved_client,
                                                           sites_file, worked):
        BW.build_worklist(saved_client, sites_path=sites_file)
        row = read_rows(worked)["Whitepages"]
        assert row["Confirmation / ref"] == "WP-88213"
        assert row["Notes"] == "phone verification, two tries"

    def test_keeps_a_partly_worked_row(self, saved_client, sites_file, worked):
        BW.build_worklist(saved_client, sites_path=sites_file)
        assert read_rows(worked)["Spokeo"]["Status"] == "verification_sent"

    def test_reports_how_much_it_kept(self, saved_client, sites_file, worked):
        stats = {}
        BW.build_worklist(saved_client, sites_path=sites_file, stats=stats)
        assert stats["preserved"] == 2

    def test_refreshes_the_catalog_guidance(self, saved_client, sites_file, worked):
        changed = [dict(s) for s in SITES]
        changed[0]["notes"] = "UPDATED GUIDANCE"
        sites_file.write_text(yaml.safe_dump(changed), encoding="utf-8")

        BW.build_worklist(saved_client, sites_path=sites_file)
        row = read_rows(worked)["Whitepages"]
        assert row["Site guidance"] == "UPDATED GUIDANCE"
        assert row["Notes"] == "phone verification, two tries"   # worker's is untouched

    def test_backs_the_previous_file_up(self, saved_client, sites_file, worked):
        stats = {}
        BW.build_worklist(saved_client, sites_path=sites_file, stats=stats)
        assert stats["backup"] and stats["backup"].is_file()

    def test_still_tracks_a_site_that_leaves_the_catalog(self, saved_client,
                                                         sites_file, worked):
        sites_file.write_text(yaml.safe_dump(SITES[1:]), encoding="utf-8")
        stats = {}
        BW.build_worklist(saved_client, sites_path=sites_file, stats=stats)

        row = read_rows(worked)["Whitepages"]
        assert stats["carried_over"] == 1
        assert row["Category"] == "(no longer in catalog)"
        assert row["Confirmation / ref"] == "WP-88213"

    def test_an_untouched_row_is_not_counted_as_work(self, saved_client, sites_file):
        BW.build_worklist(saved_client, sites_path=sites_file)
        stats = {}
        BW.build_worklist(saved_client, sites_path=sites_file, stats=stats)
        assert stats["preserved"] == 0


class TestReadTracking:
    def test_matches_a_legacy_sheet_with_no_key_column_by_site_name(self, tmp_path):
        path = tmp_path / "legacy.xlsx"
        wb = Workbook()
        ws = wb.active
        for i, h in enumerate(["Site", "Category", "Opt-out URL",
                               "Prefilled values (copy/paste)", "Supports agent?",
                               "Status", "Submitted", "Confirmation / ref",
                               "Verify by", "Notes"], start=1):
            ws.cell(row=BW.HEADER_ROW, column=i, value=h)
        ws.cell(row=BW.HEADER_ROW + 1, column=1, value="Spokeo")
        ws.cell(row=BW.HEADER_ROW + 1, column=6, value="removed")
        wb.save(path)

        assert BW.read_tracking(path)["Spokeo"]["Status"] == "removed"

    def test_a_missing_file_tracks_nothing(self, tmp_path):
        assert BW.read_tracking(tmp_path / "nope.xlsx") == {}

    def test_an_unreadable_file_tracks_nothing_rather_than_raising(self, tmp_path):
        path = tmp_path / "corrupt.xlsx"
        path.write_bytes(b"this is not a spreadsheet")
        assert BW.read_tracking(path) == {}


def test_a_worklist_open_in_excel_reports_that_clearly(saved_client, sites_file,
                                                       monkeypatch):
    BW.build_worklist(saved_client, sites_path=sites_file)

    def locked(*_args, **_kwargs):
        raise PermissionError(13, "in use")

    monkeypatch.setattr(Workbook, "save", locked)
    with pytest.raises(ClientDataError) as exc:
        BW.build_worklist(saved_client, sites_path=sites_file)
    assert "open in another program" in str(exc.value)
