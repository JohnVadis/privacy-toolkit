"""The assisted opt-out filler. It fills; it must never submit, or touch a CAPTCHA."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

import build_worklist as BW
import optout_assist as OA
from client_context import ClientDataError, build_context, output_dir


def field(**kw):
    base = {"idx": "0", "tag": "input", "type": "text", "name": "", "id": "",
            "placeholder": "", "aria": "", "autocomplete": "", "label": "",
            "options": [], "filled": False}
    base.update(kw)
    return base


@pytest.fixture
def ctx(example_client):
    return build_context(example_client)


class TestMatching:
    @pytest.mark.parametrize("kw, var", [
        ({"label": "First Name"}, "first_name"),
        ({"name": "firstName"}, "first_name"),
        ({"name": "fname"}, "first_name"),
        ({"label": "Last Name"}, "last_name"),
        ({"placeholder": "Surname"}, "last_name"),
        ({"label": "Full name"}, "full_name"),
        ({"label": "Your Name"}, "full_name"),
        ({"label": "Email Address"}, "email"),
        ({"label": "Confirm email"}, "email"),
        ({"type": "email"}, "email"),
        ({"type": "tel"}, "phone"),
        ({"label": "Street Address"}, "home_line1"),
        ({"label": "Apt / Suite"}, "home_line2"),
        ({"name": "home_city"}, "home_city"),
        ({"label": "State"}, "home_state"),
        ({"label": "ZIP Code"}, "home_zip"),
        ({"label": "Date of birth"}, "dob"),
        ({"autocomplete": "family-name", "label": "Name 2"}, "last_name"),
    ])
    def test_recognises_common_fields(self, kw, var):
        assert OA.match_variable(field(**kw)) == var

    @pytest.mark.parametrize("kw", [
        {"label": "Enter the CAPTCHA"},
        {"name": "g-recaptcha-response"},
        {"label": "Profile URL"},
        {"label": "Link to the listing"},
        {"label": "Username"},
        {"label": "Company name"},
        {"label": "Verification code"},
        {"label": "Reason for request"},
    ])
    def test_leaves_alone_what_it_must_not_or_cannot_know(self, kw):
        assert OA.match_variable(field(**kw)) is None

    # Each of these was a wrong fill on a live site (run of 2026-10-12).
    @pytest.mark.parametrize("kw, var", [
        ({"label": "Authorized Agent Email"}, "agent_email"),         # Acxiom
        ({"label": "Organizational (Legal) Entity Name"}, None),     # Acxiom
        ({"label": "Maiden Name or Alias Name:"}, None),              # MyLife
        ({"label": "First Last Name"}, "full_name"),                  # USPhoneBook
        ({"label": "Driver's License State", "tag": "select"}, None), # LexisNexis
        ({"aria": "Search for properties",
          "placeholder": "Enter your street address"}, "home_line1"), # Redfin claim
        ({"aria": "Search for properties",
          "placeholder": "City, Address, School, Agent, ZIP"}, None), # Redfin
        ({"placeholder": "City, Address, ZIP"}, None),
        ({"label": "Location"}, "home_city_state"),                   # Whitepages
        ({"placeholder": "City, State"}, "home_city_state"),
        ({"name": "month"}, "dob_month"),                             # Instant Checkmate
        ({"name": "day"}, "dob_day"),
        ({"label": "Your birth year helps us locate your correct profile"},
         "dob_year"),                                                  # MyLife
        ({"label": "Date of Birth", "placeholder": "MM/DD/YYYY"}, "dob"),  # Acxiom
        ({"label": "Postal / Zip Code"}, "home_zip"),
        ({"label": "Email Address"}, "email"),
    ])
    def test_live_site_labels(self, kw, var):
        assert OA.match_variable(field(**kw)) == var

    # Found in the audit of 2026-10-13: sites use type="tel" for any digits, and
    # a box's label beats its type.
    @pytest.mark.parametrize("kw, var", [
        ({"label": "ZIP Code", "type": "tel"}, "home_zip"),
        ({"label": "Date of Birth", "type": "tel", "placeholder": "MM/DD/YYYY"}, "dob"),
        ({"label": "Age", "type": "tel"}, None),
        ({"label": "Middle Initial"}, "middle_initial"),
        ({"name": "mi"}, "middle_initial"),
        ({"label": "Middle Name"}, "middle_name"),
        ({"label": "Prior address"}, None),
        ({"aria": "Name", "placeholder": "Enter name, phone or address"}, None),  # TPS
        ({"label": "Phone or email"}, None),
        ({"label": "Email address"}, "email"),
        ({"placeholder": "Enter your home address",
          "autocomplete": "new-password"}, "home_line1"),                  # Zillow
        ({"label": "Code", "autocomplete": "one-time-code"}, None),
        ({"label": "Address where you lived before"}, None),
    ])
    def test_label_beats_input_type(self, kw, var):
        assert OA.match_variable(field(**kw)) == var

    # Each of these would have put the client's details where they don't belong.
    @pytest.mark.parametrize("label", [
        "Place of birth", "City of birth", "Card expiration month", "Expiry year",
        "Year you moved in", "Name suffix", "Name on card", "Spouse name",
        "First name of spouse", "Parent or guardian name", "Country/Region",
        "Country", "IP address", "Web address", "Number of years at address",
        "Past addresses", "Phone type", "Phone extension", "Mobile carrier",
        "Last 4 digits of phone", "Last name, first name", "ZIP+4 extension",
    ])
    def test_other_facts_are_left_alone(self, label):
        assert OA.match_variable(field(label=label)) is None, label

    @pytest.mark.parametrize("kw, var", [
        ({"label": "Authorized representative email"}, "agent_email"),
        ({"label": "Authorized representative name"}, "agent_name"),
    ])
    def test_representative_is_the_agent(self, kw, var):
        assert OA.match_variable(field(**kw)) == var

    def test_birthday_dropdowns(self, ctx):
        month = field(tag="select", name="month", options=[
            {"value": "0", "text": "Month"}, {"value": "4", "text": "April"}])
        day = field(idx="1", tag="select", name="day", options=[
            {"value": "", "text": "Select"}, {"value": "17", "text": "17"}])
        year = field(idx="2", tag="select", name="year", options=[
            {"value": "1986", "text": "1986"}])
        fills, _ = OA.plan_fills([month, day, year], ctx)
        assert [f.value for f in fills] == ["4", "17", "1986"]

    def test_city_state_box_gets_both(self, ctx):
        (fill,), _ = OA.plan_fills([field(label="Location")], ctx)
        assert fill.value == "Brandon, FL"

    def test_never_overwrites_a_field_with_something_in_it(self, ctx):
        fills, _ = OA.plan_fills([field(label="First Name", filled=True)], ctx)
        assert fills == []

    def test_state_select_takes_the_sites_own_option(self, ctx):
        sel = field(tag="select", label="State", options=[
            {"value": "", "text": "Choose"}, {"value": "fl", "text": "Florida"}])
        (fill,), _ = OA.plan_fills([sel], ctx)
        assert (fill.kind, fill.value) == ("select", "fl")

    def test_select_without_a_matching_option_is_left_not_guessed(self, ctx):
        sel = field(tag="select", label="State", options=[{"value": "x", "text": "Texas"}])
        fills, left = OA.plan_fills([sel], ctx)
        assert fills == [] and left == [sel]

    def test_date_input_gets_iso_dob(self, ctx):
        (fill,), _ = OA.plan_fills([field(type="date", label="Date of birth")], ctx)
        assert fill.value == "1986-04-17"

    def test_middle_initial_gets_one_letter(self, ctx):
        ctx["middle_name"] = "marie"
        (fill,), _ = OA.plan_fills([field(label="Middle Initial")], ctx)
        assert fill.value == "M"

    def test_empty_client_value_is_left_for_the_worker(self, ctx):
        ctx["middle_name"] = ""
        fills, left = OA.plan_fills([field(label="Middle name")], ctx)
        assert fills == [] and len(left) == 1


def test_off_site_tabs_are_noticed():
    site = {"opt_out_url": "https://www.instantcheckmate.com/privacy-center"}
    assert not OA.off_site("https://app.instantcheckmate.com/privacy-center/", site)
    assert OA.off_site("https://ads.example.net/lead-form", site)
    assert not OA.off_site("https://anything", None)


class TestRecordProgress:
    def test_a_row_inserted_above_the_headers_is_still_found(self, saved_client):
        path = BW.build_worklist(saved_client)
        BW.record_progress(saved_client, "spokeo", status="submitted", confirmation="REF1")
        wb = load_workbook(path)
        wb.active.insert_rows(1)
        wb.save(path)
        BW.record_progress(saved_client, "whitepages", status="submitted")
        tracked = BW.read_tracking(path)
        assert tracked["spokeo"]["Confirmation / ref"] == "REF1"     # not wiped
        assert tracked["whitepages"]["Status"] == "submitted"

    def test_a_sheet_it_cant_read_is_never_rebuilt_over(self, saved_client):
        path = BW.build_worklist(saved_client)
        BW.record_progress(saved_client, "spokeo", status="submitted", confirmation="REF1")
        wb = load_workbook(path)
        ws = wb.active
        for c in ws[BW.HEADER_ROW]:
            if c.value == "Notes":
                c.value = "My notes"                   # renamed by the worker
        wb.save(path)
        before = path.read_bytes()
        with pytest.raises(ClientDataError):
            BW.record_progress(saved_client, "whitepages", status="submitted")
        assert path.read_bytes() == before

    def test_backups_in_the_same_second_dont_overwrite_each_other(self, saved_client):
        path = BW.build_worklist(saved_client)
        first = BW._backup_existing(path)
        second = BW._backup_existing(path)
        assert first != second and first.is_file() and second.is_file()

    def test_records_into_the_worklist_and_keeps_notes(self, saved_client):
        path = BW.build_worklist(saved_client)
        BW.record_progress(saved_client, "spokeo", note="first note")
        BW.record_progress(saved_client, "spokeo", status="submitted",
                           submitted="10/09/2026", confirmation="ABC123",
                           note="second")
        ws = load_workbook(path).active
        headers = [c.value for c in ws[BW.HEADER_ROW]]
        row = next(dict(zip(headers, r)) for r in
                   ws.iter_rows(min_row=BW.HEADER_ROW + 1, values_only=True)
                   if r[headers.index("Key")] == "spokeo")
        assert row["Status"] == "submitted"
        assert row["Confirmation / ref"] == "ABC123"
        assert row["Notes"] == "first note\nsecond"

    def test_builds_the_worklist_if_there_isnt_one(self, saved_client):
        path = BW.record_progress(saved_client, "whitepages", status="submitted")
        assert path.is_file()

    def test_survives_a_rebuild(self, saved_client):
        BW.record_progress(saved_client, "radaris", status="verification_sent")
        BW.build_worklist(saved_client)
        tracked = BW.read_tracking(output_dir(saved_client) / BW.WORKLIST_FILENAME)
        assert tracked["radaris"]["Status"] == "verification_sent"

    def test_refuses_unknown_site_and_status(self, saved_client):
        with pytest.raises(ClientDataError):
            BW.record_progress(saved_client, "no_such_site", status="submitted")
        with pytest.raises(ClientDataError):
            BW.record_progress(saved_client, "spokeo", status="done")

    def test_worked_sites_drop_out_of_the_default_run(self, saved_client):
        BW.record_progress(saved_client, "spokeo", status="submitted")
        keys = [s["key"] for s in OA.select_sites(saved_client, None, False, None)]
        assert "spokeo" not in keys and "whitepages" in keys
        everything = OA.select_sites(saved_client, None, True, None)
        assert "spokeo" in [s["key"] for s in everything]


# ---------------------------------------------------------------------------
# End to end, in a real (headless) browser, against a local fake opt-out form.
# No network: the page is set from a string.
# ---------------------------------------------------------------------------
FAKE_FORM = """
<form id="f" onsubmit="window.SUBMITTED = true; return false;">
  <label for="fn">First Name</label><input id="fn">
  <label for="ln">Last Name</label><input id="ln">
  <label>Email <input type="email" name="email"></label>
  <label for="city">City</label><input id="city" value="Already typed">
  <label for="st">State</label>
  <select id="st"><option value="">--</option><option value="FL">Florida</option></select>
  <label for="url">Profile URL</label><input id="url">
  <label for="cap">Type the characters (captcha)</label><input id="cap">
  <input type="hidden" name="token" value="t">
  <input type="text" id="invisible" name="first_name" style="display:none">
  <input type="text" id="offpage" name="last_name" style="position:absolute;left:-9999px">
  <input type="email" id="clear" name="email2" style="opacity:0">
  <div aria-hidden="true"><label>Phone <input id="ariahidden"></label></div>
  <input type="text" id="tiny" name="phone" style="width:1px;height:1px;padding:0;border:0">
  <button type="submit">Remove my info</button>
</form>
<script>window.SUBMITTED = false;</script>
"""


@pytest.fixture(scope="module")
def browser():
    sync_api = pytest.importorskip("playwright.sync_api")
    with sync_api.sync_playwright() as pw:
        b = None
        for channel in OA.BROWSER_CHANNELS:
            try:
                b = pw.chromium.launch(channel=channel, headless=True)
                break
            except Exception:
                continue
        if b is None:
            pytest.skip("no browser Playwright can launch")
        yield b
        b.close()


def test_fills_the_form_and_never_submits(browser, example_client):
    page = browser.new_page()
    try:
        page.set_content(FAKE_FORM)
        ctx = build_context(example_client)
        report = OA.fill_page(page, ctx)

        assert page.input_value("#fn") == "Alexandra"
        assert page.input_value("#ln") == "Example"
        assert page.input_value("input[name=email]") == "alex.example@example.com"
        assert page.input_value("#st") == "FL"
        assert page.input_value("#city") == "Already typed"      # worker's typing wins
        assert page.input_value("#url") == ""                   # the worker's call
        assert page.input_value("#cap") == ""                   # never a CAPTCHA
        assert page.input_value("#invisible") == ""             # hidden honeypot
        for honeypot in ("#offpage", "#clear", "#ariahidden", "#tiny"):
            assert page.input_value(honeypot) == "", honeypot
        assert page.evaluate("window.SUBMITTED") is False       # a human clicks
        assert {v for _, v in report["filled"]} == {
            "first_name", "last_name", "email", "home_state"}
        assert report["errors"] == []
    finally:
        page.close()


def test_a_dropdown_without_a_placeholder_is_the_workers_choice(browser, example_client):
    page = browser.new_page()
    try:
        page.set_content('<label for="st">State</label><select id="st">'
                         '<option value="AL">Alabama</option><option value="FL">Florida</option>'
                         '</select><label for="m">Month</label><select id="m">'
                         '<option value="0">Month</option><option value="4">April</option></select>'
                         '<label for="s2">Search state</label><select id="s2">'
                         '<option value="ALL">All States</option><option value="FL">Florida</option>'
                         '</select>')
        OA.fill_page(page, build_context(example_client))
        assert page.input_value("#st") == "AL"     # might be what the worker picked
        assert page.input_value("#m") == "4"       # a placeholder: fair to fill
        assert page.input_value("#s2") == "FL"     # "All States" is a placeholder too
    finally:
        page.close()


def test_stops_when_the_site_moves_on_by_itself(browser, example_client):
    page = browser.new_page()
    try:
        page.set_content('<label for="fn">First Name</label>'
                         '<input id="fn" onchange="location.hash = \'sent\'">'
                         '<label for="ln">Last Name</label><input id="ln">'
                         '<label for="em">Email</label><input id="em">')
        report = OA.fill_page(page, build_context(example_client))
        assert report["moved_on"] is True
    finally:
        page.close()


def test_a_second_fill_after_the_page_changes(browser, example_client):
    page = browser.new_page()
    try:
        page.set_content('<label for="a">First Name</label><input id="a">'
                         '<div id="later"></div>')
        OA.fill_page(page, build_context(example_client))
        page.evaluate("""() => { document.getElementById('a').style.display = 'none';
            document.getElementById('later').innerHTML =
              '<label for=b>Last Name</label><input id=b>'; }""")
        report = OA.fill_page(page, build_context(example_client))
        assert page.input_value("#b") == "Example"
        assert report["errors"] == []
    finally:
        page.close()


def test_site_selector_entries_go_first(browser, example_client):
    page = browser.new_page()
    try:
        page.set_content('<input id="weird" name="q1"><input id="q2">')
        site = {"fill": [{"selector": "#weird", "from": "home_zip"}]}
        OA.fill_page(page, build_context(example_client), site)
        assert page.input_value("#weird") == "33511"
        assert page.input_value("#q2") == ""
    finally:
        page.close()


# ---------------------------------------------------------------------------
# The window (optout_panel.py)
# ---------------------------------------------------------------------------
class TestSession:
    def test_walks_the_sites_and_records(self, saved_client):
        import optout_panel as OP

        by_key = {s["key"]: s for s in BW.load_sites()}
        sites = [by_key["spokeo"], by_key["radaris"]]
        s = OP.Session(saved_client, None, sites)
        assert s.site["key"] == "spokeo"
        assert "Alexandra M Example" in [v for _, v in s.values()]
        s.record("submitted", "REF1")
        assert s.site["key"] == "radaris"
        s.skip()
        assert s.finished
        assert s.outcomes == [("Spokeo", "submitted"), ("Radaris", "skipped")]
        assert OA.site_status(saved_client)["spokeo"] == "submitted"

    def test_a_failed_save_stays_on_the_site(self, saved_client, monkeypatch):
        import optout_panel as OP

        def boom(*a, **k):
            raise ClientDataError("open in Excel")
        monkeypatch.setattr(OP, "record_site", boom)
        s = OP.Session(saved_client, None, BW.load_sites()[:1])
        with pytest.raises(ClientDataError):
            s.record("submitted")
        assert s.index == 0 and s.outcomes == []

    def test_spouse_progress_goes_to_notes_not_status(self, saved_client):
        saved_client["persons"].append({"role": "spouse", "first_name": "Sam",
                                        "last_name": "Example"})
        ctx = build_context(saved_client, "spouse")
        site = next(s for s in BW.load_sites() if s["key"] == "spokeo")
        OA.record_site(saved_client, site, ctx, "submitted", "R9")
        tracked = BW.read_tracking(output_dir(saved_client) / BW.WORKLIST_FILENAME)
        assert "Sam Example (spouse): submitted" in tracked["spokeo"]["Notes"]
        assert tracked["spokeo"]["Status"] == "not_started"

    def test_worklist_person_is_the_first_one_whatever_the_role(self, saved_client):
        # The worklist is built for persons[0]. Here that's the spouse — so the
        # spouse's progress is the Status column, and the "primary" goes to Notes.
        saved_client["persons"].insert(0, {"role": "spouse", "first_name": "Sam",
                                           "last_name": "Example"})
        assert OA.is_worklist_person(saved_client, "spouse")
        assert not OA.is_worklist_person(saved_client, "primary")
        site = next(s for s in BW.load_sites() if s["key"] == "spokeo")
        OA.record_site(saved_client, site, build_context(saved_client, "primary"),
                       "submitted")
        tracked = BW.read_tracking(output_dir(saved_client) / BW.WORKLIST_FILENAME)
        assert tracked["spokeo"]["Status"] == "not_started"
        assert "(primary): submitted" in tracked["spokeo"]["Notes"]


# Runs in a CHILD process. Tk objects must be created and destroyed on one thread,
# and in the pytest process some were garbage-collected later on a web-app test's
# worker thread: a warning on Windows, an abort of the whole run on macOS.
WINDOW_WALK = r"""
import sys
import tkinter
import optout_assist as OA
import optout_panel as OP
import build_worklist as BW
from client_context import client_path, load_client

class FakeBrowser:
    def __init__(self, *a, **k):
        self.opened = []
    def open(self, url):
        self.opened.append(url)
        return True
    def current_page(self):
        return None
    def close(self):
        pass

OP.Browser = FakeBrowser
try:
    app = OP.App()
except tkinter.TclError:
    print("SKIP no display")
    sys.exit(0)
try:
    assert [c["slug"] for c in app.clients] == ["Example Client"], app.clients
    ticked = [s["key"] for s, var, _ in app.site_vars if var.get()]
    assert len(ticked) == len(BW.load_sites())          # nothing worked yet
    for _site, var, _status in app.site_vars:
        var.set(False)
    app.site_vars[0][1].set(True)
    app.start()
    assert app.browser.opened == [app.site_vars[0][0]["opt_out_url"]]
    app.fill()                                          # no tab: says so, no crash
    assert "browser is closed" in app.result.cget("text")
    app.ref.insert(0, "ABC")
    app.done("submitted")
    assert app.session.finished
    assert app.session.outcomes[0][1] == "submitted"
    key = app.site_vars[0][0]["key"]
    client = load_client(client_path("Example Client"))
    assert OA.site_status(client)[key] == "submitted"
finally:
    app.quit()
print("OK")
"""


def test_window_walks_setup_site_and_summary(saved_client):
    import subprocess
    import sys
    from pathlib import Path

    pytest.importorskip("tkinter")
    repo = Path(__file__).resolve().parent.parent
    run = subprocess.run([sys.executable, "-c", WINDOW_WALK], cwd=repo,
                         capture_output=True, text=True, timeout=120)
    if "SKIP no display" in run.stdout:
        pytest.skip("no display")
    assert run.returncode == 0 and "OK" in run.stdout, run.stdout + run.stderr


def test_catalog_is_well_formed():
    """Every site opens somewhere and asks only for variables the client file has."""
    from client_context import EXAMPLE_CLIENT
    known = set(build_context(EXAMPLE_CLIENT))
    sites = BW.load_sites()
    keys = [s["key"] for s in sites]
    assert len(keys) == len(set(keys))
    for s in sites:
        assert s["opt_out_url"].startswith("https://"), s["key"]
        assert set(s.get("needs", [])) <= known, s["key"]
