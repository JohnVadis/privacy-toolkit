#!/usr/bin/env python3
"""
build_worklist.py — generate a per-client opt-out worklist spreadsheet.

    python build_worklist.py --client clients/example_client.yaml

Reads sites.yaml + the client file, pre-fills the values each site asks for, and
writes output/<client_slug>/opt_out_worklist.xlsx with status-tracking columns and
a dropdown. The worker fills in status/dates as they go — that spreadsheet is the
tracker (no database needed at this volume). Zero AI, zero tokens.

Opt-out sites are PREPARED here, never submitted: they gate on CAPTCHAs and email
verification, and a human does the clicking. Nothing in this file talks to a site.
"""
from __future__ import annotations
import argparse
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from client_context import ClientDataError, build_context, load_client
from paths import OUTPUT_DIR, SITES_FILE, resolve, slugify

STATUSES = ["not_started", "submitted", "verification_sent", "verified",
            "removed", "rejected", "not_removable"]

HEADERS = ["Site", "Category", "Opt-out URL", "Prefilled values (copy/paste)",
           "Supports agent?", "Status", "Submitted", "Confirmation / ref",
           "Verify by", "Notes"]

WORKLIST_FILENAME = "opt_out_worklist.xlsx"


def load_sites(sites_path=None) -> list[dict]:
    """The opt-out catalog. Also backs the read-only sites table in the web UI."""
    with open(resolve(sites_path or SITES_FILE), encoding="utf-8") as fh:
        return yaml.safe_load(fh) or []


def prefilled(site: dict, ctx: dict) -> str:
    parts = []
    for var in site.get("needs", []):
        val = ctx.get(var, "")
        if val:
            parts.append(f"{var}: {val}")
    return "  |  ".join(parts)


def build_worklist(client: dict, person: str | None = None, sites_path=None) -> Path:
    """Write the client's opt-out worklist and return its path.

    The one entry point; the CLI below and webapp/ both call it.
    """
    ctx = build_context(client, person)
    sites = load_sites(sites_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Opt-out worklist"

    title = f"Opt-out worklist — {ctx['case_display_name'] or ctx['full_name']}"
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    ws["A2"] = ("Prefilled values come from the client file. Fill in Status/Submitted/"
                "Confirmation as you go. Re-check removals after ~2 weeks.")
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="555555")

    header_row = 4
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border

    r = header_row + 1
    for site in sites:
        row_vals = [
            site.get("name", site.get("key", "")),
            site.get("category", ""),
            site.get("opt_out_url", ""),
            prefilled(site, ctx),
            "yes" if site.get("supports_agent") else "no",
            "not_started",
            "", "", "",
            site.get("notes", ""),
        ]
        for c, v in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        r += 1

    # Status dropdown on the Status column (F), for the data rows.
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(STATUSES), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F{header_row+1}:F{r-1}")

    widths = [22, 14, 40, 46, 14, 16, 12, 20, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row+1}"

    out_dir = OUTPUT_DIR / slugify((client.get("case") or {}).get("display_name"))
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / WORKLIST_FILENAME
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="Build an opt-out worklist for a client.")
    ap.add_argument("--client", required=True)
    ap.add_argument("--person", help="person role to use (default: first/primary)")
    ap.add_argument("--sites", default=None, help="site catalog (default: sites.yaml)")
    args = ap.parse_args()

    try:
        client = load_client(args.client)
        out_path = build_worklist(client, args.person, args.sites)
    except ClientDataError as exc:
        raise SystemExit(str(exc))

    count = len(load_sites(args.sites))
    print(f"[ok] wrote {out_path}  ({count} sites)")


if __name__ == "__main__":
    main()
